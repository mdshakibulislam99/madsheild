import unittest
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from madsheild.backend.app import config
from madsheild.backend.app.api import dependencies, endpoints
from madsheild.backend.app.api.endpoints import upload
from madsheild.backend.app.main import app
from madsheild.backend.app.services import file_manager
from madsheild.backend.app.workers import tasks


class ScanApiTests(unittest.TestCase):
    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.uploads = Path(temporary.name) / "uploads"
        for patcher in (
            patch.object(file_manager, "UPLOADS_DIR", self.uploads),
            patch.multiple(config, MODEL_PATH="", MODEL_ENDPOINT="", MODEL_API_KEY=""),
            patch.object(tasks, "engine", tasks.AnalysisEngine()),
        ):
            patcher.start()
            self.addCleanup(patcher.stop)
        self.client = TestClient(app)
        self.addCleanup(self.client.close)

    def scan(self, content=b"meeting notes", filename="report.txt"):
        return self.client.post("/api/v1/scans", files={"file": (filename, content)})

    def test_unified_server_serves_frontend_assets_and_health(self):
        self.assertEqual(self.client.get("/").status_code, 200)
        self.assertIn("Mad", self.client.get("/").text)
        self.assertEqual(self.client.get("/assets/js/api-client.js").status_code, 200)
        self.assertEqual(self.client.get("/health").status_code, 200)
        self.assertEqual(self.client.get("/api/v1/capabilities").json()["model_mode"], "not_configured")

    def test_txt_scan_and_persisted_report_are_consistent(self):
        content = b"first\npowershell -enc dGVzdA=="
        response = self.scan(content)
        self.assertEqual(response.status_code, 200, response.text)
        report = response.json()
        self.assertEqual(report["status"], "completed")
        self.assertEqual(report["sha256"], sha256(content).hexdigest())
        self.assertEqual(report["verdict"], "suspicious")
        self.assertEqual(report["threats"][0]["location"]["line_number"], 2)
        self.assertIsNone(report["malware_probability"])
        self.assertIsNone(report["confidence"])
        persisted = self.client.get("/api/v1/reports/" + report["scan_id"])
        self.assertEqual(persisted.json(), report)
        self.assertTrue((self.uploads / report["scan_id"] / "report.json").is_file())

    def test_plain_text_does_not_claim_file_is_safe(self):
        report = self.scan().json()
        self.assertEqual(report["verdict"], "no_indicators_found")
        self.assertEqual(report["risk_score_type"], "heuristic_rule_score")
        self.assertTrue(any("does not prove" in value for value in report["limitations"]))

    def test_docx_scan_reports_table_coordinates(self):
        document = Document()
        document.add_table(rows=1, cols=1).cell(0, 0).text = "vssadmin delete shadows"
        buffer = BytesIO()
        document.save(buffer)
        response = self.scan(buffer.getvalue(), "report.DOCX")
        self.assertEqual(response.status_code, 200, response.text)
        location = response.json()["threats"][0]["location"]
        self.assertEqual(location["table_location"][0], {
            "table_number": 1, "row_number": 1, "cell_number": 1,
        })

    def test_identical_uploads_have_distinct_reports_and_cached_analysis(self):
        first = self.scan().json()
        second = self.scan().json()
        self.assertNotEqual(first["scan_id"], second["scan_id"])
        self.assertFalse(first["analysis"]["cache_hit"])
        self.assertTrue(second["analysis"]["cache_hit"])

    def test_invalid_uploads_do_not_create_storage(self):
        cases = [(b"", "empty.txt", 400), (b"binary", "file.exe", 400),
                 (b"\xff", "bad.txt", 422), (b"not docx", "bad.docx", 422)]
        for content, filename, status in cases:
            with self.subTest(filename=filename):
                response = self.scan(content, filename)
                self.assertEqual(response.status_code, status, response.text)
        self.assertFalse(self.uploads.exists())

    def test_upload_size_limit_and_busy_response(self):
        with patch.object(config, "MAX_FILE_SIZE_BYTES", 3):
            self.assertEqual(self.scan(b"four").status_code, 413)
        acquired = []
        try:
            while dependencies.scan_slots.acquire(blocking=False):
                acquired.append(True)
            response = self.scan()
            self.assertEqual(response.status_code, 503)
            self.assertEqual(response.headers["retry-after"], "2")
        finally:
            for _ in acquired:
                dependencies.scan_slots.release()
        self.assertEqual(self.scan().status_code, 200)

    def test_model_failure_preserves_rule_analysis(self):
        with patch.object(config, "MODEL_PATH", "/missing/model.json"):
            response = self.scan(b"powershell -enc test")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["model"]["status"], "failed")
        self.assertEqual(response.json()["verdict"], "suspicious")
        self.assertIsNone(response.json()["malware_probability"])

    def test_invalid_or_missing_reports_are_not_found(self):
        for scan_id in ("invalid", "0" * 32):
            self.assertEqual(self.client.get("/api/v1/reports/" + scan_id).status_code, 404)


if __name__ == "__main__":
    unittest.main()
