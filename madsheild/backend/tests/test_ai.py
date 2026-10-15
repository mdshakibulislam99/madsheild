import json
import math
import unittest
import struct
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile

import httpx
from docx import Document
from pypdf import PdfWriter

from madsheild.backend.app import config
from madsheild.backend.app.ai import feature_extractor
from madsheild.backend.app.ai.feature_extractor import extract_features, extract_txt_lines
from madsheild.backend.app.ai.models.classifier import Classifier, rule_analysis
from madsheild.backend.app.workers.tasks import AnalysisEngine


def docx_bytes(document):
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pdf_bytes(*, encrypted=False):
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    if encrypted:
        writer.encrypt("test-password")
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def pe_bytes():
    content = bytearray(0x80)
    content[:2] = b"MZ"
    struct.pack_into("<I", content, 0x3C, 0x80)
    content.extend(b"PE\0\0")
    content.extend(struct.pack("<HHIIIHH", 0x14C, 0, 0, 0, 0, 0xE0, 0x0102))
    content.extend(struct.pack("<H", 0x10B) + bytes(222))
    return bytes(content)


class FeatureExtractionTests(unittest.TestCase):
    def test_text_preserves_blank_lines_and_handles_bom(self):
        self.assertEqual(extract_txt_lines(b"\xef\xbb\xbffirst\r\n\r\nthird"), [
            {"line_number": 1, "text": "first"},
            {"line_number": 2, "text": ""},
            {"line_number": 3, "text": "third"},
        ])

    def test_text_rejects_invalid_encoding_and_binary_data(self):
        for content in (b"\xff", b"hello\x00world"):
            with self.subTest(content=content), self.assertRaises(ValueError):
                extract_txt_lines(content)

    def test_additional_text_formats_preserve_line_locations(self):
        result = extract_features(b'{"command": "powershell -enc dGVzdA=="}', ".json")
        self.assertEqual(result["segments"][0]["line_number"], 1)
        self.assertEqual(result["features"]["powershell_count"], 1)

    def test_limits_prevent_unbounded_extraction(self):
        with patch.object(feature_extractor, "MAX_TEXT_CHARACTERS", 3), self.assertRaises(ValueError):
            extract_txt_lines(b"four")
        with patch.object(feature_extractor, "MAX_TEXT_SEGMENTS", 2), self.assertRaises(ValueError):
            extract_txt_lines(b"one\ntwo\nthree")

    def test_docx_preserves_blank_paragraphs_and_nested_table_location(self):
        document = Document()
        for text in ("first", "", "third"):
            document.add_paragraph(text)
        table = document.add_table(rows=2, cols=2)
        table.cell(1, 1).text = "table text"
        nested = table.cell(1, 1).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "nested text"
        result = extract_features(docx_bytes(document), ".docx")
        segments = result["segments"]
        self.assertEqual([s["paragraph_number"] for s in segments[:3]], [1, 2, 3])
        self.assertEqual(segments[1]["text"], "")
        nested_segment = next(s for s in segments if s["text"] == "nested text")
        self.assertEqual(nested_segment["table_location"], [
            {"table_number": 1, "row_number": 2, "cell_number": 2},
            {"table_number": 1, "row_number": 1, "cell_number": 1},
        ])

    def test_docx_rejects_non_zip_duplicate_entries_and_entity_declarations(self):
        with self.assertRaises(ValueError):
            extract_features(b"not a document", ".docx")
        buffer = BytesIO()
        with ZipFile(buffer, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr("word/document.xml", '<!DOCTYPE doc [<!ENTITY x "bad">]><doc/>')
        with self.assertRaises(ValueError):
            extract_features(buffer.getvalue(), ".docx")

    def test_docx_archive_indicators_are_reported(self):
        original = docx_bytes(Document())
        buffer = BytesIO()
        with ZipFile(BytesIO(original)) as source, ZipFile(buffer, "w") as target:
            for item in source.infolist():
                target.writestr(item, source.read(item.filename))
            target.writestr("word/embeddings/object.bin", b"test placeholder")
            target.writestr("word/vbaProject.bin", b"test placeholder")
        features = extract_features(buffer.getvalue(), ".docx")
        result = rule_analysis(features["segments"], features["features"])
        self.assertEqual(result["risk_score"], 100)
        self.assertEqual(len(result["findings"]), 2)
        self.assertTrue(all(f["location"]["scope"] == "document_archive" for f in result["findings"]))

    def test_pdf_extracts_page_locations_and_reports_structure_markers(self):
        result = extract_features(pdf_bytes() + b"\n/JavaScript", ".pdf")
        self.assertEqual(result["segments"][0]["page_number"], 1)
        self.assertEqual(result["features"]["pdf_javascript_count"], 1)
        findings = rule_analysis(result["segments"], result["features"])["findings"]
        self.assertEqual(findings[0]["location"]["scope"], "file_structure")

    def test_pdf_rejects_encrypted_documents(self):
        with self.assertRaises(ValueError):
            extract_features(pdf_bytes(encrypted=True), ".pdf")

    def test_zip_scans_supported_members_and_rejects_unsafe_paths(self):
        buffer = BytesIO()
        with ZipFile(buffer, "w") as archive:
            archive.writestr("notes/alert.log", b"powershell -enc dGVzdA==")
            archive.writestr("ignored.bin", b"not analyzed")
        result = extract_features(buffer.getvalue(), ".zip")
        self.assertEqual(result["features"]["archive_member_count"], 1)
        self.assertEqual(result["segments"][0]["archive_member"], "notes/alert.log")
        self.assertEqual(rule_analysis(result["segments"], result["features"])["risk_score"], 70)

        unsafe = BytesIO()
        with ZipFile(unsafe, "w") as archive:
            archive.writestr("../outside.txt", b"ordinary notes")
        with self.assertRaises(ValueError):
            extract_features(unsafe.getvalue(), ".zip")

    def test_pe_is_parsed_statically_and_rejects_invalid_content(self):
        result = extract_features(pe_bytes(), ".exe")
        self.assertEqual(result["features"]["pe_executable"], 1)
        with self.assertRaises(ValueError):
            extract_features(b"not an executable", ".exe")

    def test_legacy_doc_rejects_non_ole_content(self):
        with self.assertRaises(ValueError):
            extract_features(b"not a legacy document", ".doc")


class RuleAnalysisTests(unittest.TestCase):
    def analyze(self, content):
        features = extract_features(content, ".txt")
        return rule_analysis(features["segments"], features["features"])

    def test_rules_locate_encoded_command_without_executing_it(self):
        result = self.analyze(b"ordinary text\npowershell -enc dGVzdA==\n")
        finding = result["findings"][0]
        self.assertEqual(finding["location"]["line_number"], 2)
        self.assertEqual(finding["rule_id"], "encoded_powershell")
        self.assertEqual(result["risk_score"], 70)

    def test_repeated_rule_does_not_inflate_score_and_findings_are_bounded(self):
        result = self.analyze(b"powershell -enc dGVzdA==\n" * 105)
        self.assertEqual(result["risk_score"], 70)
        self.assertEqual(result["match_count"], 105)
        self.assertEqual(len(result["findings"]), 100)
        self.assertTrue(result["findings_truncated"])

    def test_plain_text_has_no_rule_indicators(self):
        self.assertEqual(self.analyze(b"Project meeting notes")["risk_score"], 0)


class ModelInferenceTests(unittest.TestCase):
    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.artifact = Path(temporary.name) / "model.json"
        settings = patch.multiple(config, MODEL_PATH="", MODEL_ENDPOINT="", MODEL_NAME="test",
                                  MODEL_API_KEY="", MODEL_THRESHOLD=0.8,
                                  MODEL_MAX_SEGMENTS=64, MODEL_MAX_TEXT_CHARACTERS=64000)
        settings.start()
        self.addCleanup(settings.stop)
        self.input = extract_features(b"hello", ".txt")

    def predict(self, classifier=None):
        return (classifier or Classifier()).predict(self.input["segments"], self.input["features"])

    def test_missing_model_keeps_probability_unavailable(self):
        result = self.predict()
        self.assertEqual(result["status"], "not_configured")
        self.assertIsNone(result["malware_probability"])

    def test_local_inference_applies_normalization_and_sigmoid(self):
        # Synthetic parameters verify inference mechanics; they are not a trained detector.
        self.artifact.write_text(json.dumps({
            "format": "madshield-logistic-v1", "feature_names": ["text_length"],
            "weights": [2], "means": [3], "scales": [2], "bias": -1,
        }), encoding="utf-8")
        with patch.object(config, "MODEL_PATH", str(self.artifact)):
            classifier = Classifier()
            result = self.predict(classifier)
            self.assertEqual(result["status"], "completed")
            self.assertAlmostEqual(result["malware_probability"], 1 / (1 + math.exp(-1)))
            loaded = classifier._artifact
            self.predict(classifier)
            self.assertIs(classifier._artifact, loaded)

    def test_invalid_artifact_fails_without_fabricating_probability(self):
        for artifact in ({"format": "unknown"}, {"format": "madshield-logistic-v1",
                         "feature_names": ["text_length"], "weights": [1], "scales": [0]}):
            self.artifact.write_text(json.dumps(artifact), encoding="utf-8")
            with patch.object(config, "MODEL_PATH", str(self.artifact)):
                result = self.predict()
                self.assertEqual(result["status"], "failed")
                self.assertIsNone(result["malware_probability"])

    def service(self, handler):
        client = httpx.Client(transport=httpx.MockTransport(handler))
        self.addCleanup(client.close)
        return patch("madsheild.backend.app.ai.models.classifier.httpx.Client", return_value=client)

    def test_service_prediction_preserves_location_and_discloses_partial_coverage(self):
        self.input = extract_features(b"first\nsecond", ".txt")
        def handler(request):
            payload = json.loads(request.content)
            self.assertEqual(payload["segments"][0]["location"]["line_number"], 1)
            return httpx.Response(200, json={"predictions": [
                {"id": s["id"], "malware_probability": 0.9, "reason": "test prediction"}
                for s in payload["segments"]
            ]})
        with patch.multiple(config, MODEL_ENDPOINT="http://model.test/predict", MODEL_MAX_SEGMENTS=1), self.service(handler):
            result = self.predict()
        self.assertEqual(result["status"], "completed")
        self.assertFalse(result["coverage"]["complete"])
        self.assertEqual(result["findings"][0]["location"]["line_number"], 1)
        self.assertEqual(result["malware_probability"], 0.9)

    def test_service_rejects_missing_duplicate_and_out_of_range_predictions(self):
        responses = [[], [{"id": "0", "malware_probability": 2}],
                     [{"id": "0", "malware_probability": True}],
                     [{"id": "different", "malware_probability": 0.9}]]
        for predictions in responses:
            with self.subTest(predictions=predictions):
                def handler(request):
                    return httpx.Response(200, json={"predictions": predictions})
                with patch.object(config, "MODEL_ENDPOINT", "http://model.test/predict"), self.service(handler):
                    result = self.predict()
                self.assertEqual(result["status"], "failed")
                self.assertIsNone(result["malware_probability"])


class AnalysisCacheTests(unittest.TestCase):
    def setUp(self):
        settings = patch.multiple(config, MODEL_PATH="", MODEL_ENDPOINT="", MODEL_API_KEY="")
        settings.start()
        self.addCleanup(settings.stop)

    def test_cache_reuses_extraction_and_does_not_share_mutable_results(self):
        engine = AnalysisEngine()
        first = engine.analyze(".txt", b"powershell -enc abc")
        first["rules"]["findings"].clear()
        with patch("madsheild.backend.app.workers.tasks.extract_features", side_effect=AssertionError("cache missed")):
            second = engine.analyze(".txt", b"powershell -enc abc")
        self.assertTrue(second["cache_hit"])
        self.assertEqual(len(second["rules"]["findings"]), 1)

    def test_failed_model_prediction_is_retried(self):
        engine = AnalysisEngine()
        failed = {"status": "failed", "mode": "service", "malware_probability": None, "findings": []}
        with patch.object(engine.classifier, "predict", return_value=failed) as predictor:
            engine.analyze(".txt", b"hello")
            engine.analyze(".txt", b"hello")
        self.assertEqual(predictor.call_count, 2)


if __name__ == "__main__":
    unittest.main()
